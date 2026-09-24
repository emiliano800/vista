"""Import processors. API handlers only ever talk to `ImportProcessor`; the
deterministic processor is the default and `AgentImportProcessor`
(VISTA_IMPORT_PROCESSOR=agent) layers a model call on top of it for the columns
the alias tables cannot place. Both write through the same import API to the
same canonical tables."""

from __future__ import annotations

import csv
import io
import json
import re
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from datetime import datetime

from fastapi import HTTPException

from vista.agents.llm import Prompt, chat, strip_fences
from vista.config import settings

# ---- Canonical dataset schema (the import contract every surface uses) ------------

DATASETS: dict[str, dict] = {
    "customers": {
        "label": "Customers",
        "entity": "customer",
        "fields": {
            "customer_name": {
                "label": "Customer name",
                "required": True,
                "aliases": ["client", "customer", "name", "account name", "company", "customer name", "client name", "bill to", "dba name"],
            },
            "source_customer_id": {
                "label": "Source customer ID",
                "aliases": ["customer id", "client id", "cust id", "id", "customer #", "account id", "account #"],
            },
            "contact": {"label": "Contact", "aliases": ["contact", "attn", "primary contact", "contact name"]},
            "email": {"label": "Email", "aliases": ["email", "e-mail", "contact email"]},
            "phone": {"label": "Phone", "aliases": ["phone", "tel", "telephone", "phone #", "mobile"]},
            "address": {
                "label": "Address",
                "aliases": [
                    "address",
                    "street",
                    "addr",
                    "service address",
                    "billing street",
                    "bill to street",
                    "address 1",
                    "billing address",
                ],
            },
            "city": {"label": "City", "aliases": ["city", "town", "city/state", "city, state", "location", "billing city", "bill to city"]},
            "state": {"label": "State", "aliases": ["state", "st", "province", "billing state", "bill to state"]},
            "zip": {"label": "ZIP", "aliases": ["zip", "zip code", "postal", "postal code", "billing zip", "bill to zip"]},
            "service_type": {
                "label": "Service type / segment",
                "aliases": ["service", "service type", "type", "segment", "line", "industry", "lines of business", "naics code"],
            },
            "status": {"label": "Status", "aliases": ["status", "active", "active?"]},
        },
    },
    "invoices": {
        "label": "Invoices / Accounts Receivable",
        "entity": "invoice",
        "fields": {
            "source_invoice_number": {
                "label": "Source invoice number",
                "required": True,
                "aliases": ["inv #", "invoice", "invoice #", "invoice number", "inv no", "number", "ref", "invoice id", "num", "inv"],
            },
            "customer_name": {
                "label": "Customer name",
                "aliases": ["client", "customer", "customer name", "client name", "bill to", "account"],
            },
            "source_customer_id": {"label": "Source customer ID", "aliases": ["customer id", "client id", "cust id", "account id"]},
            "issue_date": {"label": "Issue date", "aliases": ["date", "invoice date", "issued", "inv date"]},
            "due_date": {"label": "Due date", "aliases": ["due", "due date", "pay by", "payable by"]},
            "amount": {
                "label": "Amount",
                "aliases": ["amount", "total", "invoice total", "amt", "billed", "total amount", "premium", "prem", "gross"],
            },
            "outstanding_balance": {
                "label": "Outstanding balance",
                "aliases": ["balance", "amt due", "amount due", "open", "outstanding", "remaining", "open bal", "open balance", "bal"],
            },
            "status": {"label": "Status", "aliases": ["status", "paid?", "state"]},
        },
    },
    "vendors": {
        "label": "Vendors / Purchases",
        "entity": "vendor_purchase",
        "fields": {
            "vendor_name": {
                "label": "Vendor",
                "required": True,
                "aliases": ["vendor", "supplier", "payee", "vendor name", "supplier name"],
            },
            "sku": {"label": "SKU / item", "aliases": ["sku", "item", "item #", "part", "part #", "product code", "item id"]},
            "description": {"label": "Description", "aliases": ["description", "desc", "item description", "memo"]},
            "quantity": {"label": "Quantity", "aliases": ["qty", "quantity", "units"]},
            "unit": {"label": "Unit", "aliases": ["unit", "uom", "u/m"]},
            "unit_price": {"label": "Unit price", "aliases": ["price", "unit price", "rate", "cost", "unit cost"]},
            "date": {"label": "Date", "aliases": ["date", "po date", "purchase date", "ordered"]},
            "total": {"label": "Total", "aliases": ["total", "ext price", "extended", "line total", "amount"]},
        },
    },
    "vendor_invoices": {
        "label": "Vendor invoices / Accounts Payable",
        "entity": "vendor_purchase",
        "fields": {
            "vendor_name": {
                "label": "Vendor",
                "required": True,
                "aliases": ["vendor name", "vendor", "supplier", "payee", "supplier name"],
            },
            "source_vendor_id": {"label": "Source vendor ID", "aliases": ["vendor id", "supplier id", "payee id"]},
            "invoice_number": {
                "label": "Invoice number",
                "aliases": [
                    "invoice number",
                    "invoice #",
                    "inv #",
                    "invoice",
                    "ap invoice id",
                    "supplier invoice number",
                    "supplier invoice n",
                    "vendor invoice number",
                ],
            },
            "date": {"label": "Invoice date", "aliases": ["invoice date", "date", "inv date"]},
            "due_date": {"label": "Due date", "aliases": ["due date", "due"]},
            "total": {"label": "Amount", "aliases": ["amount", "total", "amt", "invoice total"]},
            "category": {"label": "Category", "aliases": ["category", "expense category", "type"]},
            "gl_account": {"label": "GL account", "aliases": ["gl account", "gl", "account"]},
            "status": {"label": "Status", "aliases": ["status", "paid?"]},
        },
    },
    "vendor_master": {
        "label": "Vendor / supplier master",
        "entity": "vendor",
        "fields": {
            "vendor_name": {
                "label": "Vendor",
                "required": True,
                "aliases": ["vendor name", "vendor", "supplier name", "supplier", "payee", "name"],
            },
            "source_vendor_id": {"label": "Source vendor ID", "aliases": ["vendor id", "supplier id", "payee id", "id"]},
            "category": {"label": "Category", "aliases": ["category", "type", "commodity"]},
            "payment_terms": {"label": "Payment terms", "aliases": ["payment terms", "terms"]},
            "contact": {"label": "Contact", "aliases": ["contact", "contact name", "remit to"]},
            "email": {"label": "Email", "aliases": ["email", "e-mail"]},
            "phone": {"label": "Phone", "aliases": ["phone", "tel"]},
            "city": {"label": "City", "aliases": ["city"]},
            "state": {"label": "State", "aliases": ["state", "st"]},
            "status": {"label": "Status", "aliases": ["status", "active", "active?", "approved supplier"]},
        },
    },
    "subscriptions": {
        "label": "Software subscriptions",
        "entity": "subscription",
        "fields": {
            "product": {
                "label": "Product",
                "required": True,
                "aliases": ["product", "software", "application", "tool", "service", "product name", "subscription"],
            },
            "vendor": {"label": "Vendor", "aliases": ["vendor", "publisher", "provider", "vendor name"]},
            "category": {"label": "Category", "aliases": ["category", "type", "function"]},
            "monthly_cost": {"label": "Monthly cost", "aliases": ["monthly $", "monthly", "monthly cost", "price", "mrr", "cost"]},
            "annual_cost": {"label": "Annual cost", "aliases": ["annual cost", "annual $", "annual", "yearly cost", "arr"]},
            "seats": {"label": "Seats", "aliases": ["seats", "users", "licenses", "licences", "seats licensed"]},
            "renewal_date": {"label": "Renewal date", "aliases": ["renewal", "renews", "renewal date", "term end", "expires"]},
            "notes": {"label": "Contract notes", "aliases": ["notes", "terms", "contract", "restrictions", "contract term"]},
        },
    },
    "policies": {
        "label": "Insurance policies",
        "entity": "policy",
        "fields": {
            "policy_number": {
                "label": "Policy number",
                "required": True,
                "aliases": ["policy number", "policy #", "policy no", "pol #", "pol no"],
            },
            "source_policy_id": {"label": "Source policy ID", "aliases": ["policy id", "pol id"]},
            "customer_name": {
                "label": "Client name",
                "aliases": ["client name", "client", "customer name", "insured", "named insured", "customer"],
            },
            "source_customer_id": {"label": "Source client ID", "aliases": ["client id", "customer id", "cust id", "insured id"]},
            "line_of_business": {"label": "Line of business", "aliases": ["line of business", "lob", "line", "coverage line"]},
            "line_description": {"label": "Line description", "aliases": ["line description", "coverage", "coverage description"]},
            "carrier_code": {"label": "Carrier code", "aliases": ["carrier code", "carrier id", "carrier"]},
            "carrier_name": {"label": "Carrier name", "aliases": ["carrier name", "insurer", "company"]},
            "effective_date": {"label": "Effective date", "aliases": ["effective date", "eff date", "effective", "eff", "inception"]},
            "expiration_date": {
                "label": "Expiration date",
                "aliases": ["expiration date", "exp date", "expiration", "exp", "expires", "expiry"],
            },
            "term_months": {"label": "Term (months)", "aliases": ["term months", "term", "term mo"]},
            "annual_premium": {
                "label": "Annual premium",
                "aliases": ["annual premium", "premium", "prem", "written premium", "total premium"],
            },
            "commission_pct": {
                "label": "Commission %",
                "aliases": ["commission pct", "commission %", "comm pct", "comm %", "commission rate"],
            },
            "expected_commission": {
                "label": "Expected commission",
                "aliases": ["expected commission", "comm est", "commission", "est commission", "commission amount", "comm"],
            },
            "billing_type": {"label": "Billing type", "aliases": ["billing type", "bill type", "billing"]},
            "status": {"label": "Status", "aliases": ["policy status", "status"]},
            "producer_id": {"label": "Producer", "aliases": ["producer id", "producer"]},
            "account_manager_id": {"label": "Account manager", "aliases": ["account manager id", "account manager", "am id", "csr"]},
            "surplus_lines": {"label": "Surplus lines", "aliases": ["surplus lines", "surplus", "e&s"]},
            "new_or_renewal": {"label": "New / renewal", "aliases": ["new or renewal", "new/renewal", "transaction type", "new renewal"]},
            "experience_mod": {"label": "Experience mod", "aliases": ["experience mod", "exp mod", "x mod", "emod"]},
            "umbrella_limit": {"label": "Umbrella limit", "aliases": ["umbrella limit", "umbrella"]},
        },
    },
    "purchase_orders": {
        "label": "Purchase orders",
        "entity": "purchase_order",
        "fields": {
            "po_number": {
                "label": "PO number",
                "required": True,
                "aliases": ["po number", "po #", "po", "po no", "purchase order", "po num"],
            },
            "vendor_name": {"label": "Supplier", "aliases": ["supplier name", "supplier", "vendor name", "vendor", "payee"]},
            "source_vendor_id": {"label": "Source supplier ID", "aliases": ["supplier id", "vendor id"]},
            "date": {"label": "PO date", "aliases": ["po date", "date", "order date", "ordered"]},
            "buyer_id": {"label": "Buyer", "aliases": ["buyer id", "buyer"]},
            "payment_terms": {"label": "Payment terms", "aliases": ["payment terms", "terms"]},
            "ship_via": {"label": "Ship via", "aliases": ["ship via", "carrier", "shipping method"]},
            "freight_terms": {"label": "Freight terms", "aliases": ["freight terms", "freight", "fob"]},
            "total": {"label": "PO total", "aliases": ["po total", "total", "amount", "order total"]},
            "status": {"label": "Status", "aliases": ["po status", "status"]},
            "approved_by": {"label": "Approved by", "aliases": ["approved by", "approver"]},
            "sent_method": {"label": "Sent method", "aliases": ["sent method", "sent via", "delivery method"]},
        },
    },
    "purchase_order_lines": {
        "label": "Purchase order lines",
        "entity": "purchase_order_line",
        "fields": {
            "po_number": {
                "label": "PO number",
                "required": True,
                "aliases": ["po number", "po #", "po", "po no", "purchase order", "po num"],
            },
            "line_number": {"label": "Line", "aliases": ["line", "line #", "line number", "line no"]},
            "item_id": {"label": "Item", "aliases": ["item id", "item", "sku", "part", "part #", "item #"]},
            "description": {"label": "Description", "aliases": ["description", "desc", "item description"]},
            "manufacturer_part_number": {
                "label": "Manufacturer part number",
                "aliases": ["manufacturer part number", "manufacturer part", "mpn", "mfr part", "mfg part number", "mfr part #"],
            },
            "quantity": {"label": "Ordered qty", "aliases": ["ordered qty", "qty", "quantity", "qty ordered", "order qty"]},
            "unit": {"label": "UOM", "aliases": ["uom", "unit", "u/m"]},
            "unit_price": {"label": "Unit cost", "aliases": ["unit cost", "unit price", "price", "cost", "rate"]},
            "total": {
                "label": "Extended cost",
                "aliases": ["extended cost", "extended", "ext cost", "line total", "extended price", "amount"],
            },
            "need_by_date": {"label": "Need-by date", "aliases": ["need by date", "need by", "required date", "requested date"]},
            "promised_date": {"label": "Promised date", "aliases": ["promised date", "promised", "due date"]},
            "received_qty": {"label": "Received qty", "aliases": ["received qty", "received", "qty received"]},
            "status": {"label": "Line status", "aliases": ["line status", "status"]},
            "gl_account": {"label": "GL account", "aliases": ["gl account", "gl", "account"]},
        },
    },
    "inventory": {
        "label": "Inventory balances",
        "entity": "inventory_balance",
        "fields": {
            "item_id": {"label": "Item", "required": True, "aliases": ["item id", "item", "sku", "part", "part #", "item #"]},
            "warehouse": {"label": "Warehouse", "aliases": ["warehouse", "site", "location", "plant"]},
            "bin_location": {"label": "Bin", "aliases": ["bin location", "bin", "bin loc"]},
            "on_hand_qty": {"label": "On hand", "aliases": ["on hand qty", "on hand", "qty on hand", "onhand", "oh qty"]},
            "allocated_qty": {"label": "Allocated", "aliases": ["allocated qty", "allocated", "committed", "reserved"]},
            "available_qty": {"label": "Available", "aliases": ["available qty", "available", "avail"]},
            "on_order_qty": {"label": "On order", "aliases": ["on order qty", "on order", "on po"]},
            "unit": {"label": "UOM", "aliases": ["uom", "unit", "u/m"]},
            "unit_price": {"label": "Unit cost", "aliases": ["unit cost", "cost", "std cost", "unit price"]},
            "total": {"label": "Extended value", "aliases": ["extended value", "ext value", "value", "extended cost", "total value"]},
            "last_count_date": {"label": "Last count", "aliases": ["last count date", "last count", "counted"]},
            "last_receipt_date": {"label": "Last receipt", "aliases": ["last receipt date", "last receipt", "last received"]},
            "last_issue_date": {"label": "Last issue", "aliases": ["last issue date", "last issue", "last used"]},
            "as_of_date": {"label": "As of", "aliases": ["as of date", "as of", "snapshot date"]},
        },
    },
    "profile": {"label": "Company profile", "entity": None, "fields": {}},
    "other": {"label": "Other", "entity": None, "fields": {}},
}
MAPPABLE = [k for k, d in DATASETS.items() if d["fields"]]

# Confidence at or above which a proposed column mapping is applied without a human decision.
AUTO_CONFIDENCE = 0.9


@dataclass
class Inspection:
    columns: list[str]
    rows: list[dict[str, str]]
    sheet: str
    dataset: str
    confidence: float


@dataclass
class ProposedMapping:
    source: str
    example: str
    target: str | None
    confidence: float
    reason: str = ""

    @property
    def needs_review(self) -> bool:
        return not (self.target and self.confidence >= AUTO_CONFIDENCE)


@dataclass
class NormalizedRecord:
    source_row: int
    raw: dict
    normalized: dict
    confidence: float
    # Cells that could not be read as the number their field requires: {field, column, value}.
    problems: list[dict] = field(default_factory=list)


@dataclass
class DetectedException:
    exception_type: str
    description: str
    left: str
    right: str
    confidence: float
    actions: list[str]
    dataset: str
    left_row: int | None = None
    right_row: int | None = None
    record_rows: list[int] = field(default_factory=list)
    match_vendor: str | None = None
    candidates: list[dict] = field(default_factory=list)
    field_name: str | None = None  # the normalized field an "Unreadable amount" refers to


@dataclass
class ModelCall:
    """One model consultation made by a processor, so the caller can meter it on an AgentRun."""

    purpose: str
    model: str
    input_tokens: int
    output_tokens: int
    source: str
    columns: list[str]


class ImportProcessor(ABC):
    name = "abstract"

    @abstractmethod
    def inspect_file(self, filename: str, content: bytes, sheet: str | None = None) -> Inspection: ...

    @abstractmethod
    def propose_mappings(self, dataset: str, columns: list[str], rows: list[dict]) -> list[ProposedMapping]: ...

    @abstractmethod
    def normalize_records(self, dataset: str, rows: list[dict], mappings: list[ProposedMapping]) -> list[NormalizedRecord]: ...

    @abstractmethod
    def identify_exceptions(
        self, dataset: str, records: list[NormalizedRecord], existing_vendors: list[str]
    ) -> list[DetectedException]: ...

    def list_sheets(self, filename: str, content: bytes) -> list[str]:
        """Worksheet names for workbooks; a single pseudo-sheet for flat files."""
        if filename.lower().endswith((".xlsx", ".xlsm")):
            return workbook_sheets(content)
        return ["Sheet1"]


# ---- Parsing --------------------------------------------------------------------


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9#$/?&]+", " ", str(s).lower()).strip()


def parse_csv(text: str) -> tuple[list[str], list[dict[str, str]]]:
    rows = [r for r in csv.reader(io.StringIO(text)) if any(c.strip() for c in r)]
    if not rows:
        return [], []
    header = [h.strip() for h in rows[0]]
    body = [{h: (r[i].strip() if i < len(r) else "") for i, h in enumerate(header)} for r in rows[1:]]
    return header, body


def workbook_sheets(content: bytes) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    return [ws.title for ws in wb.worksheets]


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if (v.hour, v.minute, v.second) == (0, 0, 0) else v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_xlsx(content: bytes, sheet: str | None = None) -> tuple[str, list[str], list[dict[str, str]]]:
    """One worksheet (the first when `sheet` is None) as header + rows of strings."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    if sheet is not None:
        if sheet not in wb.sheetnames:
            raise HTTPException(422, f"Worksheet {sheet!r} not found; sheets are {', '.join(wb.sheetnames)}")
        ws = wb[sheet]
    else:
        ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it, [])]
    body = []
    for r in it:
        if not any(v not in (None, "") for v in r):
            continue
        body.append({h: _cell(v) for h, v in zip(header, r, strict=False) if h})
    return ws.title, [h for h in header if h], body


# ---- Normalisation ------------------------------------------------------------------

STATES = {
    "alabama": "AL", "alaska": "AK", "arizona": "AZ", "arkansas": "AR", "california": "CA", "colorado": "CO", "connecticut": "CT",
    "delaware": "DE", "florida": "FL", "georgia": "GA", "hawaii": "HI", "idaho": "ID", "illinois": "IL", "indiana": "IN", "iowa": "IA",
    "kansas": "KS", "kentucky": "KY", "louisiana": "LA", "maine": "ME", "maryland": "MD", "massachusetts": "MA", "mass": "MA",
    "michigan": "MI", "minnesota": "MN", "mississippi": "MS", "missouri": "MO", "montana": "MT", "nebraska": "NE", "nevada": "NV",
    "new hampshire": "NH", "new jersey": "NJ", "new mexico": "NM", "new york": "NY", "north carolina": "NC", "north dakota": "ND",
    "ohio": "OH", "oklahoma": "OK", "oregon": "OR", "pennsylvania": "PA", "rhode island": "RI", "south carolina": "SC",
    "south dakota": "SD", "tennessee": "TN", "texas": "TX", "utah": "UT", "vermont": "VT", "virginia": "VA", "washington": "WA",
    "west virginia": "WV", "wisconsin": "WI", "wyoming": "WY",
}  # fmt: skip


def normalize_text(v) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip()


def normalize_state(v) -> str:
    s = str(v or "").strip().rstrip(".")
    if re.fullmatch(r"[A-Za-z]{2}", s):
        return s.upper()
    return STATES.get(s.lower(), s)


def normalize_city_state(v) -> tuple[str, str]:
    s = normalize_text(v)
    m = re.match(r"^(.*?)[,\s]+([A-Za-z.]{2,14})$", s)
    if not m:
        return s, ""
    state = normalize_state(m.group(2))
    if not re.fullmatch(r"[A-Z]{2}", state):
        return s, ""
    return m.group(1).strip().rstrip(","), state


def normalize_phone(v) -> str:
    digits = re.sub(r"\D", "", str(v or ""))
    if len(digits) == 11 and digits.startswith("1"):
        digits = digits[1:]
    if len(digits) != 10:
        return str(v or "").strip()
    return f"({digits[:3]}) {digits[3:6]}-{digits[6:]}"


_DATE_FORMATS = ("%b %d, %Y", "%B %d, %Y", "%d %b %Y", "%d-%b-%Y", "%d-%b-%y", "%d %B %Y", "%Y/%m/%d", "%b %d %Y", "%d.%m.%Y")


def normalize_date(v) -> str:
    s = str(v or "").strip()
    if not s:
        return ""
    if re.match(r"^\d{4}-\d{2}-\d{2}", s):
        return s[:10]
    m = re.fullmatch(r"(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})", s)
    if m:
        year = f"20{m.group(3)}" if len(m.group(3)) == 2 else m.group(3)
        return f"{year}-{m.group(1).zfill(2)}-{m.group(2).zfill(2)}"
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return s


_MONEY_RE = re.compile(
    r"^\(?\s*[-+]?\s*(?:[$€£]|USD|EUR|GBP|CAD)?\s*[-+]?\s*(?:\d[\d,]*(?:\.\d*)?|\.\d+)\s*%?\s*(?:USD|EUR|GBP|CAD)?\s*\)?\s*-?$",
    re.IGNORECASE,
)


def normalize_money(v) -> float | None:
    """A number, 0.0 for an empty cell, or None when the cell holds text that is not a
    number ("N/A", "TBD", "see note"). None is an import exception, never a silent zero."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    s = str(v or "").strip()
    if not s:
        return 0.0
    if not _MONEY_RE.match(s):
        return None
    negative = (s.startswith("(") and s.endswith(")")) or s.lstrip("($€£ ").startswith("-") or s.endswith("-")
    n = float(re.sub(r"[^0-9.]", "", s))
    return -n if negative else n


def normalize_bool(v) -> bool:
    return str(v or "").strip().lower() in ("y", "yes", "true", "1", "x")


_STATUS_MAP = {
    "y": "active", "yes": "active", "true": "active", "1": "active", "active": "active",
    "n": "inactive", "no": "inactive", "false": "inactive", "0": "inactive", "inactive": "inactive",
    "open": "open", "closed": "closed", "cancelled": "cancelled", "canceled": "cancelled", "void": "cancelled",
    "paid": "paid", "pd": "paid", "settled": "paid", "partial": "partial", "partially paid": "partial",
    "overdue": "overdue", "late": "overdue", "past due": "overdue",
    "in force": "in_force", "inforce": "in_force", "bound": "in_force", "expired": "expired", "non renewed": "non_renewed",
    "received": "received", "partially received": "partial", "sent": "open", "approved": "open", "draft": "draft",
    "pending": "pending", "prospect": "prospect",
}  # fmt: skip


def normalize_status(v) -> str:
    s = re.sub(r"\s+", " ", str(v or "").strip().lower())
    if s in _STATUS_MAP:
        return _STATUS_MAP[s]
    if s.startswith("past due") or s.startswith("overdue"):
        return "overdue"
    return s


def normalize_name(v) -> str:
    s = normalize_text(v)
    s = re.sub(r"\b(l\.?l\.?c\.?|inc\.?|co\.?|corp\.?|ltd\.?)\b", lambda m: m.group(0).replace(".", "").upper(), s, flags=re.I)
    return re.sub(r"\s*,\s*", ", ", s)


DATE_FIELDS = {
    "issue_date", "due_date", "date", "renewal_date", "effective_date", "expiration_date",
    "need_by_date", "promised_date", "last_count_date", "last_receipt_date", "last_issue_date", "as_of_date",
}  # fmt: skip
MONEY_FIELDS = {
    "amount", "outstanding_balance", "unit_price", "total", "monthly_cost", "annual_cost",
    "annual_premium", "expected_commission", "umbrella_limit",
}  # fmt: skip
DECIMAL_FIELDS = {
    "quantity", "received_qty", "on_hand_qty", "allocated_qty", "available_qty", "on_order_qty",
    "commission_pct", "experience_mod",
}  # fmt: skip
COUNT_FIELDS = {"seats", "line_number", "term_months"}
# A numeric cell that holds text. The row is rejected at approval unless the analyst decides.
UNREADABLE_AMOUNT = "Unreadable amount"
UNREADABLE_AMOUNT_ACTIONS = ["Skip row", "Import as zero"]
BOOL_FIELDS = {"surplus_lines"}
NAME_FIELDS = {"customer_name", "vendor_name", "carrier_name"}
UPPER_FIELDS = {"unit", "line_of_business", "carrier_code"}


def _name_key(s: str) -> str:
    k = normalize_text(s).lower()
    k = re.sub(r"[.,']", "", k)
    k = re.sub(r"\b(llc|inc|co|corp|ltd|the|services?|enterprises?|supply|hvac|boston|w w)\b", "", k)
    k = re.sub(r"#\d+", "", k)
    return re.sub(r"\s+", " ", k).strip()


def similarity(a: str, b: str) -> float:
    x, y = _name_key(a), _name_key(b)
    if not x or not y:
        return 0.0
    if x == y:
        return 0.93
    wa, wb = set(x.split(" ")), set(y.split(" "))
    j = len(wa & wb) / len(wa | wb)
    return 0.6 + j * 0.3 if j >= 0.5 else 0.0


# ---- Deterministic processor ------------------------------------------------------

_DETECT_BONUS = {
    "invoices": r"\bar\b|invoice|receivable",
    "customers": r"customer|client|account",
    "vendors": r"purchase|\bpo\b|vendor_purchase",
    "vendor_invoices": r"ap_|payable|vendor_invoice|ap vendor",
    "vendor_master": r"vendors?\.|suppliers?\.|corporate_vendors|vendor master|supplier master",
    "subscriptions": r"software|subscription|saas|license",
    "policies": r"polic",
    "purchase_orders": r"purchase_orders?\.|\bpos?\b|purchase orders?\.",
    "purchase_order_lines": r"po_lines|order_lines|po lines|lines",
    "inventory": r"inventory|stock|on_hand|balances",
}
_SHARED_HEADS = {"date", "total", "amount", "status", "type", "id", "name", "category", "description", "account", "terms"}


class DeterministicDemoImportProcessor(ImportProcessor):
    """Header-alias mapping, rule-based normalisation and similarity-based
    exception detection. No model calls; every output is reproducible."""

    name = "demo"

    def inspect_file(self, filename: str, content: bytes, sheet: str | None = None) -> Inspection:
        lower = filename.lower()
        if lower.endswith((".xlsx", ".xlsm")):
            sheet_name, columns, rows = parse_xlsx(content, sheet)
        elif lower.endswith((".csv", ".txt")):
            sheet_name = "Sheet1"
            columns, rows = parse_csv(content.decode("utf-8-sig", errors="replace"))
        else:
            raise HTTPException(415, f"{filename}: only CSV and XLSX exports are parsed by the {self.name} import processor")
        if not columns:
            raise HTTPException(422, f"{filename}: no header row found")
        hint = f"{filename} {sheet_name}" if sheet_name != "Sheet1" else filename
        dataset, confidence = self.detect_dataset(columns, hint)
        return Inspection(columns, rows, sheet_name, dataset if confidence >= 0.5 else "other", confidence)

    @staticmethod
    def detect_dataset(columns: list[str], filename: str = "") -> tuple[str, float]:
        heads = [_norm(c) for c in columns]
        name = filename.lower()
        scores = []
        for key in MAPPABLE:
            fields = DATASETS[key]["fields"]
            hits = 0.0
            for h in heads:
                for f in fields.values():
                    if h in f["aliases"]:
                        # Exact, dataset-specific headers outweigh shared ones like "date"/"status".
                        hits += 0.5 if h in _SHARED_HEADS else 1.0
                        break
            required_hit = all(any(h in f["aliases"] for h in heads) for f in fields.values() if f.get("required"))
            bonus = 0.15 if re.search(_DETECT_BONUS[key], name) else 0.0
            score = (hits / len(heads) + bonus + (0.1 if required_hit else -0.2)) if heads else 0.0
            scores.append((key, score))
        scores.sort(key=lambda kv: -kv[1])
        best, runner = scores[0], scores[1] if len(scores) > 1 else (None, 0.0)
        margin = best[1] - runner[1]
        confidence = max(0.0, min(0.99, best[1] * 0.75 + margin * 0.6 + 0.1))
        return best[0], round(confidence, 2)

    def propose_mappings(self, dataset: str, columns: list[str], rows: list[dict]) -> list[ProposedMapping]:
        fields = DATASETS.get(dataset, DATASETS["other"])["fields"]
        # Score every (column, field) pair, then assign greedily from the strongest
        # match down so a weak alias never steals a field from a later exact header.
        candidates: list[tuple[float, int, str, str]] = []
        for ci, column in enumerate(columns):
            head = _norm(column)
            for key, d in fields.items():
                aliases = d["aliases"]
                idx = aliases.index(head) if head in aliases else -1
                if idx == 0 or head == _norm(d["label"]) or head == _norm(key):
                    score = 0.99
                elif 0 < idx <= 3:
                    score = 0.97 - idx * 0.001
                elif idx > 3:
                    score = 0.93 - idx * 0.001
                elif len(head) >= 3 and any(a in head or head in a for a in aliases):
                    score = 0.72
                else:
                    continue
                candidates.append((score, ci, column, key))
        candidates.sort(key=lambda c: (-c[0], c[1]))
        chosen: dict[str, tuple[str, float]] = {}
        taken: set[str] = set()
        for score, _ci, column, key in candidates:
            if column in chosen or key in taken:
                continue
            chosen[column] = (key, round(score, 2))
            taken.add(key)
        out = []
        for column in columns:
            target, confidence = chosen.get(column, (None, 0.0))
            if not target and re.match(r"^acct|^acc|account", _norm(column)):
                confidence = 0.51
            example = next((r[column] for r in rows if r.get(column)), "")
            out.append(ProposedMapping(column, example, target, confidence, "header alias" if target else ""))
        return out

    def normalize_records(self, dataset: str, rows: list[dict], mappings: list[ProposedMapping]) -> list[NormalizedRecord]:
        fields = DATASETS.get(dataset, DATASETS["other"])["fields"]
        active = [m for m in mappings if m.target and m.target in fields]
        confidence = min([m.confidence for m in active] + [1.0])
        out = []
        for i, row in enumerate(rows):
            record: dict = {}
            raw: dict = {}
            problems: list[dict] = []
            for m in active:
                value = row.get(m.source, "")
                raw[m.source] = value
                t = m.target
                if t in MONEY_FIELDS or t in DECIMAL_FIELDS or t in COUNT_FIELDS:
                    n = normalize_money(value)
                    if n is None:
                        problems.append({"field": t, "column": m.source, "value": str(value)})
                        record[t] = None
                        continue
                    record[t] = round(n) if t in COUNT_FIELDS else n
                    continue
                if t == "phone":
                    v = normalize_phone(value)
                elif t == "state":
                    v = normalize_state(value)
                elif t == "city":
                    v, st = normalize_city_state(value)
                    if st and not record.get("state"):
                        record["state"] = st
                elif t in DATE_FIELDS:
                    v = normalize_date(value)
                elif t in BOOL_FIELDS:
                    v = normalize_bool(value)
                elif t == "status":
                    v = normalize_status(value)
                elif t in NAME_FIELDS:
                    v = normalize_name(value)
                elif t in UPPER_FIELDS:
                    v = normalize_text(value).upper()
                else:
                    v = normalize_text(value)
                record[t] = v
            out.append(NormalizedRecord(i + 2, raw, record, min(confidence, 0.5) if problems else confidence, problems))
        return out

    def identify_exceptions(self, dataset: str, records: list[NormalizedRecord], existing_vendors: list[str]) -> list[DetectedException]:
        out: list[DetectedException] = []
        for r in records:
            for p in r.problems:
                label = DATASETS.get(dataset, DATASETS["other"])["fields"].get(p["field"], {}).get("label", p["field"])
                out.append(
                    DetectedException(
                        UNREADABLE_AMOUNT,
                        f"Row {r.source_row}: {label} is '{p['value']}' (column '{p['column']}'), which is not a number. "
                        "Importing it as zero would understate the figure.",
                        p["value"],
                        label,
                        1.0,
                        UNREADABLE_AMOUNT_ACTIONS,
                        dataset,
                        left_row=r.source_row,
                        record_rows=[r.source_row],
                        field_name=p["field"],
                    )
                )
        if dataset == "customers":
            names = [(r.source_row, r.normalized.get("customer_name", "")) for r in records]
            for i in range(len(names)):
                for j in range(i + 1, len(names)):
                    s = similarity(names[i][1], names[j][1])
                    if s >= 0.85 and names[i][1] != names[j][1]:
                        out.append(
                            DetectedException(
                                "Possible duplicate customer",
                                f"'{names[i][1]}' (row {names[i][0]}) and '{names[j][1]}' (row {names[j][0]}) look like the same customer.",
                                names[i][1],
                                names[j][1],
                                round(s, 2),
                                ["Merge", "Keep separate", "Skip"],
                                "customers",
                                left_row=names[i][0],
                                right_row=names[j][0],
                                record_rows=[names[i][0], names[j][0]],
                            )
                        )
                        if len(out) >= 8:
                            return out
        if dataset in ("vendors", "vendor_invoices", "vendor_master", "purchase_orders"):
            seen: list[str] = []
            for r in records:
                v = r.normalized.get("vendor_name", "")
                if not v or v in seen:
                    continue
                seen.append(v)
                ranked = sorted(((e, similarity(v, e)) for e in existing_vendors), key=lambda es: -es[1])
                if not ranked:
                    continue
                best, s = ranked[0]
                candidates = [{"name": e, "similarity": round(sc, 2)} for e, sc in ranked[:3] if sc > 0]
                if s >= 0.6 and _name_key(v) != _name_key(best) and s < 0.93:
                    out.append(
                        DetectedException(
                            "Possible vendor match",
                            f"'{v}' may be portfolio vendor '{best}'.",
                            v,
                            best,
                            round(s, 2),
                            ["Match", "Keep separate", "Skip"],
                            "vendors",
                            match_vendor=best,
                            candidates=candidates,
                        )
                    )
                elif _name_key(v) == _name_key(best) and v != best:
                    out.append(
                        DetectedException(
                            "Vendor name variant",
                            f"'{v}' is a spelling variant of portfolio vendor '{best}'.",
                            v,
                            best,
                            0.9,
                            ["Match", "Keep separate", "Skip"],
                            "vendors",
                            match_vendor=best,
                            candidates=candidates,
                        )
                    )
        if dataset == "invoices":
            seen_numbers: dict[str, int] = {}
            for r in records:
                n = str(r.normalized.get("source_invoice_number") or "")
                if n and n in seen_numbers and len(out) < 8:
                    out.append(
                        DetectedException(
                            "Duplicate invoice number",
                            f"Invoice {n} appears on rows {seen_numbers[n]} and {r.source_row}.",
                            n,
                            n,
                            0.95,
                            ["Keep both", "Skip duplicate"],
                            "invoices",
                            left_row=seen_numbers[n],
                            right_row=r.source_row,
                            record_rows=[seen_numbers[n], r.source_row],
                        )
                    )
                seen_numbers.setdefault(n, r.source_row)
        if dataset == "policies":
            inverted = [
                r.source_row
                for r in records
                if r.normalized.get("effective_date")
                and r.normalized.get("expiration_date")
                and str(r.normalized["expiration_date"]) < str(r.normalized["effective_date"])
            ]
            if inverted:
                share = len(inverted) / max(len(records), 1)
                out.append(
                    DetectedException(
                        "Expiration precedes effective date",
                        f"{len(inverted)} of {len(records)} policies expire before they take effect"
                        + (" — the effective/expiration columns are probably swapped in the source." if share > 0.5 else "."),
                        "expiration_date",
                        "effective_date",
                        round(0.6 + 0.35 * share, 2),
                        ["Swap dates", "Keep as-is", "Skip"],
                        "policies",
                        left_row=inverted[0],
                        record_rows=inverted,
                    )
                )
        if dataset == "inventory":
            off = [
                r.source_row
                for r in records
                if r.normalized.get("available_qty") not in (None, "")
                and abs(
                    float(r.normalized.get("on_hand_qty") or 0)
                    - float(r.normalized.get("allocated_qty") or 0)
                    - float(r.normalized["available_qty"])
                )
                > 0.001
            ]
            if off:
                out.append(
                    DetectedException(
                        "Inventory quantities do not reconcile",
                        f"On hand − allocated ≠ available on {len(off)} of {len(records)} rows.",
                        "on_hand_qty − allocated_qty",
                        "available_qty",
                        0.9,
                        ["Keep as-is", "Skip"],
                        "inventory",
                        left_row=off[0],
                        record_rows=off,
                    )
                )
        return out


# ---- Agent processor: deterministic first, model only for the leftovers -----------

_MAPPING_SYSTEM = (
    "You map spreadsheet column headers to a fixed canonical schema for a private-equity data room. "
    "You receive the dataset, the canonical fields that are still unmapped, and the source columns that "
    "deterministic alias matching could not place, each with example values. Return JSON only: "
    '{"mappings": [{"source": <column>, "target": <field key or null>, "confidence": <0-1>, "reason": <short>}]}. '
    "Map a column only when the header AND the example values support it; otherwise target null. "
    "Never map two columns to the same field. Do not invent fields."
)


class AgentImportProcessor(DeterministicDemoImportProcessor):
    """Same contract as the deterministic processor. Columns it leaves unmapped
    or below `AUTO_CONFIDENCE` are sent to the model together with the fields
    still open; the answer is capped below the auto-apply threshold so a human
    (or the loader acting as one) still confirms it. Every call is recorded in
    `model_calls` for the caller to meter."""

    name = "agent"

    def __init__(self) -> None:
        self.model_calls: list[ModelCall] = []

    def propose_mappings(self, dataset: str, columns: list[str], rows: list[dict]) -> list[ProposedMapping]:
        proposals = super().propose_mappings(dataset, columns, rows)
        fields = DATASETS.get(dataset, DATASETS["other"])["fields"]
        taken = {p.target for p in proposals if p.target and p.confidence >= AUTO_CONFIDENCE}
        open_fields = {k: d for k, d in fields.items() if k not in taken}
        unresolved = [p for p in proposals if p.needs_review]
        if not open_fields or not unresolved:
            return proposals
        examples = {p.source: [str(r.get(p.source, ""))[:60] for r in rows[:40] if r.get(p.source)][:3] for p in unresolved}
        user = json.dumps(
            {
                "dataset": dataset,
                "open_fields": [
                    {"key": k, "label": d["label"], "aliases": d["aliases"][:6], "required": bool(d.get("required"))}
                    for k, d in open_fields.items()
                ],
                "columns": [
                    {"source": p.source, "examples": examples[p.source], "deterministic_guess": p.target, "guess_confidence": p.confidence}
                    for p in unresolved
                ],
            },
            ensure_ascii=False,
        )
        result = chat(Prompt(system=_MAPPING_SYSTEM, user=user, max_tokens=600))
        self.model_calls.append(
            ModelCall(
                "column_mapping", result.model, result.input_tokens, result.output_tokens, result.source, [p.source for p in unresolved]
            )
        )
        answers = self._parse(result.text)
        by_source = {p.source: p for p in proposals}
        assigned = set(taken)
        for a in answers:
            p = by_source.get(str(a.get("source", "")))
            target = a.get("target")
            if p is None or not p.needs_review or target not in open_fields or target in assigned:
                continue
            try:
                conf = float(a.get("confidence", 0.0))
            except (TypeError, ValueError):
                continue
            if conf < 0.5:
                continue
            p.target = target
            # Model proposals never auto-apply: cap under the human-review threshold.
            p.confidence = min(round(conf, 2), AUTO_CONFIDENCE - 0.01)
            p.reason = f"model: {str(a.get('reason', ''))[:160]}"
            assigned.add(target)
        return proposals

    @staticmethod
    def _parse(text: str) -> list[dict]:
        if not text.strip():
            return []
        try:
            data = json.loads(strip_fences(text))
        except json.JSONDecodeError:
            return []
        items = data.get("mappings") if isinstance(data, dict) else data
        return [i for i in items if isinstance(i, dict)] if isinstance(items, list) else []


def get_processor(name: str | None = None) -> ImportProcessor:
    name = name or settings.import_processor
    if name == "agent" and settings.enable_agent_import:
        return AgentImportProcessor()
    return DeterministicDemoImportProcessor()
