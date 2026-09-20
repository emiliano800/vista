"""File ingestion and deterministic, source-linked insurance checks.

No model is trusted with arithmetic or record matching. Original columns and row
numbers survive normalization; ambiguous keys are never silently joined.
"""

import base64
import binascii
import csv
import hashlib
import io
import re
import zipfile
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import PurePosixPath

from openpyxl import load_workbook

MAX_FILES = 12
MAX_ROWS = 1999
MAX_COLUMNS = 64
MAX_BYTES = 5 * 1024 * 1024

# Required fields drive classification. Optional fields enrich the evidence.
SCHEMAS = {
    "clients": {"label": "Clients", "required": ["client_id", "client_name"], "optional": ["status"]},
    "policies": {
        "label": "Policies",
        "required": ["policy_number", "commission_pct"],
        "optional": ["client_name", "carrier_code", "annual_premium", "policy_id"],
    },
    "commissions": {
        "label": "Commission statements",
        "required": ["policy_number", "premium_basis", "commission_paid"],
        "optional": ["statement_id", "carrier_code", "insured_name"],
    },
    "carriers": {
        "label": "Carrier agreements",
        "required": ["carrier_code", "carrier_name", "commission_pct"],
        "optional": [],
    },
    "invoices": {
        "label": "Invoices",
        "required": ["invoice_id", "balance", "due_date"],
        "optional": ["client_id", "policy_number", "payment_plan", "status"],
    },
}
ALIASES = {
    "policy_number": ["policy_no", "policy_num", "policy_#"],
    "commission_pct": ["commission_rate", "contracted_rate", "commission_percent"],
    "premium_basis": ["statement_premium", "premium_amount"],
    "commission_paid": ["paid_commission", "commission_amount"],
    "client_id": ["customer_id", "account_id"],
    "client_name": ["customer_name", "insured_name"],
    "invoice_id": ["invoice_number", "invoice_no"],
    "balance": ["outstanding_balance", "amount_outstanding"],
    "due_date": ["payment_due_date"],
}


def normalized(value):
    return re.sub(r"[\s-]+", "_", value.strip().lower())


def suggest_mapping(columns, kind):
    schema = SCHEMAS[kind]
    result = {}
    for field in schema["required"] + schema["optional"]:
        matches = [c for c in columns if normalized(c) in [field, *ALIASES.get(field, [])]]
        if len(matches) == 1:
            result[field] = matches[0]
    return result


def table_from_rows(filename, sheet, rows, index):
    if not rows or len(rows) < 2:
        raise ValueError(f"{filename}: add a header and at least one record.")
    columns = [str(c).strip() for c in rows[0]]
    if not columns or len(columns) > MAX_COLUMNS or any(not c for c in columns) or len(set(columns)) != len(columns):
        raise ValueError(f"{filename}: use unique, non-empty headers (up to {MAX_COLUMNS} columns).")
    records = []
    for row_number, row in enumerate(rows[1:], 2):
        if not any(str(c).strip() for c in row):
            continue
        if len(row) > len(columns):
            raise ValueError(f"{filename}, row {row_number}: more values than column headers.")
        if any(len(str(c)) > 512 for c in row):
            raise ValueError(f"{filename}, row {row_number}: a cell exceeds 512 characters.")
        records.append(
            {"row": row_number, "values": dict(zip(columns, [str(c) for c in row] + [""] * (len(columns) - len(row)), strict=True))}
        )
    if not records or len(records) > MAX_ROWS:
        raise ValueError(f"{filename}: use between 1 and {MAX_ROWS:,} records per table.")
    candidates = []
    for kind, schema in SCHEMAS.items():
        mapping = suggest_mapping(columns, kind)
        if all(f in mapping for f in schema["required"]):
            candidates.append(kind)
    # Policy exports also contain carrier names/rates; their policy identifier
    # makes them policy records rather than standalone carrier agreements.
    if "policies" in candidates:
        candidates = [c for c in candidates if c not in ("carriers", "clients")]
    kind = candidates[0] if len(candidates) == 1 else "unclassified"
    return {
        "id": str(index),
        "filename": filename,
        "sheet": sheet,
        "columns": columns,
        "records": records,
        "kind": kind,
        "mapping": suggest_mapping(columns, kind) if kind in SCHEMAS else {},
    }


def parse_files(files):
    if not 1 <= len(files) <= MAX_FILES:
        raise ValueError(f"Choose between 1 and {MAX_FILES} files.")
    tables, size, names = [], 0, set()
    for file in files:
        name = PurePosixPath(file["name"].replace("\\", "/")).name
        if not name or len(name) > 180 or name in names:
            raise ValueError("Each file needs a unique filename of at most 180 characters.")
        names.add(name)
        try:
            raw = base64.b64decode(file["content"], validate=True)
        except (ValueError, binascii.Error) as exc:
            raise ValueError(f"{name}: invalid file encoding.") from exc
        size += len(raw)
        if size > MAX_BYTES:
            raise ValueError("Choose files totaling at most 5 MiB.")
        ext = PurePosixPath(name).suffix.lower()
        if ext in (".csv", ".tsv"):
            try:
                reader = csv.reader(io.StringIO(raw.decode("utf-8-sig")), delimiter="\t" if ext == ".tsv" else ",", strict=True)
                rows = []
                for row in reader:
                    rows.append(row)
                    if len(rows) > MAX_ROWS + 1:
                        raise ValueError(f"{name}: maximum {MAX_ROWS:,} records per table.")
            except (UnicodeDecodeError, csv.Error) as exc:
                raise ValueError(f"{name}: use a valid UTF-8 CSV or TSV export.") from exc
            tables.append(table_from_rows(name, "", rows, len(tables)))
        elif ext == ".xlsx":
            for sheet, rows in workbook_rows(name, raw):
                tables.append(table_from_rows(name, sheet, rows, len(tables)))
        else:
            raise ValueError(f"{name}: supported formats are CSV, TSV and XLSX.")
    if len(tables) > 20:
        raise ValueError("Choose at most 20 tables across all files.")
    return tables


def workbook_rows(name, raw):
    """Read data exports strictly: no silent cell truncation or formula evaluation."""
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 20 * 1024 * 1024:
                raise ValueError(f"{name}: workbook expands beyond 20 MiB. Split the export.")
            for part in archive.namelist():
                if part.endswith(".xml"):
                    data = archive.read(part)
                    if b"<!DOCTYPE" in data or b"<!ENTITY" in data:
                        raise ValueError(f"{name}: workbook contains unsupported XML declarations.")
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            if len(workbook.worksheets) > 20:
                raise ValueError(f"{name}: maximum 20 sheets per import.")
            result = []
            for sheet in workbook.worksheets:
                if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > MAX_COLUMNS:
                    raise ValueError(f"{name}: sheet exceeds {MAX_ROWS:,} records or {MAX_COLUMNS} columns.")
                rows = []
                for cells in sheet.iter_rows():
                    if len(rows) >= MAX_ROWS + 1 or len(cells) > MAX_COLUMNS:
                        raise ValueError(f"{name}: sheet exceeds the import limits.")
                    row = []
                    for cell in cells:
                        if cell.data_type in ("f", "e"):
                            raise ValueError(f"{name}, {sheet.title}!{cell.coordinate}: export formulas/errors as values before importing.")
                        value = cell.value
                        if isinstance(value, (date, datetime)):
                            value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
                        elif isinstance(value, (float, int)) and "%" in cell.number_format:
                            value = str(Decimal(str(value)) * 100) + "%"
                        row.append("" if value is None else str(value))
                    rows.append(row)
                if any(any(c for c in row) for row in rows):
                    result.append((sheet.title, rows))
            if not result:
                raise ValueError(f"{name}: workbook has no records.")
            return result
        finally:
            workbook.close()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"{name}: could not read this workbook. Export it as CSV and try again.") from exc


def money(value):
    try:
        amount = Decimal(str(value).replace(",", "").replace("$", "").replace("%", "").strip())
        if not amount.is_finite() or abs(amount) > Decimal("1000000000000"):
            raise ValueError("Invalid amount")
        return amount
    except (InvalidOperation, ValueError) as exc:
        raise ValueError("Use a valid finite amount.") from exc


def apply_mappings(tables, choices):
    choices = {c["id"]: c for c in choices}
    if set(choices) != {t["id"] for t in tables}:
        raise ValueError("Review the mapping for every table.")
    for table in tables:
        choice = choices[table["id"]]
        kind, mapping = choice["kind"], choice["mapping"]
        if kind not in SCHEMAS:
            raise ValueError(f"{table['filename']}: choose a supported record type.")
        schema = SCHEMAS[kind]
        fields = schema["required"] + schema["optional"]
        if any(f not in mapping for f in schema["required"]) or any(f not in fields for f in mapping):
            raise ValueError(f"{table['filename']}: map every required field.")
        if any(c not in table["columns"] for c in mapping.values()) or len(set(mapping.values())) != len(mapping):
            raise ValueError(f"{table['filename']}: map each source column only once.")
        table["kind"], table["mapping"] = kind, mapping
        for record in table["records"]:
            values = {field: record["values"][column].strip() for field, column in mapping.items()}
            for field in schema["required"]:
                if not values[field]:
                    raise ValueError(f"{table['filename']}, row {record['row']}: {field} is empty.")
            for field in ("commission_pct", "premium_basis", "commission_paid", "balance", "annual_premium"):
                if values.get(field):
                    try:
                        amount = money(values[field])
                        if field == "commission_pct" and not 0 <= amount <= 100:
                            raise ValueError("Invalid rate")
                    except ValueError as exc:
                        raise ValueError(f"{table['filename']}, row {record['row']}: invalid {field}.") from exc
            if values.get("due_date"):
                try:
                    date.fromisoformat(values["due_date"])
                except ValueError as exc:
                    raise ValueError(f"{table['filename']}, row {record['row']}: use YYYY-MM-DD for due_date.") from exc
    return tables


def analyze(tables, as_of):
    grouped = {kind: [] for kind in SCHEMAS}
    for table in tables:
        for record in table["records"]:
            values = {f: record["values"][c].strip() for f, c in table["mapping"].items()}
            evidence = {
                "table_id": table["id"],
                "filename": table["filename"],
                "sheet": table["sheet"],
                "row": record["row"],
                "values": record["values"],
            }
            grouped[table["kind"]].append((values, evidence))
    findings = []

    def add(category, title, detail, recommendation, evidence, amount=None, calculation=None):
        identity = f"{category}:{[(e['table_id'], e['row']) for e in evidence]}"
        findings.append(
            {
                "id": hashlib.sha256(identity.encode()).hexdigest()[:16],
                "category": category,
                "kind": "observed_fact",
                "title": title,
                "detail": detail,
                "recommendation": recommendation,
                "evidence": evidence,
                "amount": str(amount) if amount is not None else None,
                "calculation": calculation,
                "status": "open",
            }
        )

    # Duplicate keys cannot be safely joined; surface them rather than taking the last row.
    policies = {}
    for values, evidence in grouped["policies"]:
        policies.setdefault(values["policy_number"].casefold(), []).append((values, evidence))
    for matches in policies.values():
        if len(matches) > 1:
            add(
                "data_quality",
                "Policy number appears more than once",
                f"{matches[0][0]['policy_number']} has {len(matches)} records.",
                "Confirm which policy term applies before reconciling commissions.",
                [e for _, e in matches],
            )
    matched = 0
    statement_groups = {}
    for statement, source in grouped["commissions"]:
        fingerprint = tuple(sorted(statement.items()))
        statement_groups.setdefault(fingerprint, []).append((statement, source))
    duplicates = {key for key, records in statement_groups.items() if len(records) > 1}
    for key in duplicates:
        records = statement_groups[key]
        add(
            "data_quality",
            "Identical commission records need review",
            records[0][0]["policy_number"],
            "Confirm whether these are duplicates or separate transactions. They are excluded from the discrepancy total.",
            [source for _, source in records],
        )
    for statement, source in grouped["commissions"]:
        if tuple(sorted(statement.items())) in duplicates:
            continue
        matches = policies.get(statement["policy_number"].casefold(), [])
        if not matches:
            if grouped["policies"]:
                add(
                    "data_quality",
                    "Commission has no matching policy",
                    statement["policy_number"],
                    "Check the policy export and carrier statement before recognizing this commission.",
                    [source],
                )
            continue
        if len(matches) != 1:
            continue
        policy, policy_source = matches[0]
        if statement.get("carrier_code") and policy.get("carrier_code") and statement["carrier_code"] != policy["carrier_code"]:
            add(
                "data_quality",
                "Carrier differs between matched records",
                statement["policy_number"],
                "Confirm the carrier before reconciling this statement.",
                [source, policy_source],
            )
            continue
        matched += 1
        premium, rate, paid = money(statement["premium_basis"]), money(policy["commission_pct"]), money(statement["commission_paid"])
        expected = (premium * rate / 100).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        gap = expected - paid
        if gap > 50:
            insured = statement.get("insured_name") or policy.get("client_name") or statement["policy_number"]
            add(
                "commission",
                "Commission below the policy rate",
                f"{insured} · {statement['policy_number']}",
                "Ask the carrier to reconcile the rate and confirm any adjustments before requesting a correction.",
                [source, policy_source],
                gap,
                {"premium": str(premium), "rate": str(rate), "expected": str(expected), "paid": str(paid), "difference": str(gap)},
            )
    # A payment plan changes what is collectible: only flag full-pay invoices with a supplied plan.
    for invoice, source in grouped["invoices"]:
        age = (as_of - date.fromisoformat(invoice["due_date"])).days
        if age > 60 and money(invoice["balance"]) > 0 and invoice.get("payment_plan", "").casefold() == "full pay":
            add(
                "receivable",
                "Full-pay invoice is more than 60 days overdue",
                f"{invoice['invoice_id']} · {age} days past its due date as of {as_of.isoformat()}.",
                "Confirm subsequent payments and contact the account owner. The balance is receivable exposure, not savings.",
                [source],
                money(invoice["balance"]),
            )
    recovery = sum((Decimal(f["amount"]) for f in findings if f["category"] == "commission"), Decimal(0))
    kinds = {t["kind"] for t in tables}
    checks = [
        {
            "name": "Commission reconciliation",
            "status": "completed" if {"policies", "commissions"} <= kinds else "needs_data",
            "detail": f"{matched} statement rows matched to a unique policy."
            if {"policies", "commissions"} <= kinds
            else "Import policies and commission statements together.",
        },
        {
            "name": "Receivables review",
            "status": "completed" if "invoices" in kinds else "needs_data",
            "detail": "Checks full-pay invoices over 60 days; payment-plan balances require separate review."
            if "invoices" in kinds
            else "Import invoices with due dates, balances and payment plans.",
        },
        {
            "name": "Policy identity",
            "status": "completed" if "policies" in kinds else "needs_data",
            "detail": "Duplicate policy numbers are flagged and excluded from commission matching."
            if "policies" in kinds
            else "Import policies to check for ambiguous identifiers.",
        },
    ]
    return {
        "findings": findings,
        "checks": checks,
        "summary": {
            "records": sum(len(t["records"]) for t in tables),
            "tables": len(tables),
            "clients": len(grouped["clients"]),
            "policies": len(grouped["policies"]),
            "commission_variance": str(recovery),
            "findings": len(findings),
            "matched_commissions": matched,
        },
    }
