"""Sanity checks over generated output.

    python3 synthetic_data/generator/validate.py

Checks: no empty folders, every high-tier FK resolves, planted anomalies are
actually present in the files, answer-key evidence paths exist (globs allowed),
and tier degradation happened (dropped datasets absent, workbook present).
"""
from __future__ import annotations

import csv
import glob
import json
import os
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.abspath(os.path.join(HERE, ".."))
errors: list[str] = []


def err(msg: str) -> None:
    errors.append(msg)


def load(path: str) -> list[dict]:
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def fk(child: str, ccol: str, parent: str, pcol: str) -> None:
    kids = load(child)
    parents = {r[pcol] for r in load(parent)}
    missing = [r[ccol] for r in kids if r[ccol] and r[ccol] not in parents]
    if missing:
        err(f"FK {os.path.relpath(child, OUT)}.{ccol} -> {os.path.basename(parent)}.{pcol}: {len(missing)} dangling e.g. {missing[:3]}")


def check_layout() -> None:
    for root, dirs, files in os.walk(OUT):
        if "generator" in root:
            continue
        if not dirs and not files:
            err(f"empty folder {root}")
    for f in glob.glob(os.path.join(OUT, "**", "*.csv"), recursive=True):
        rows = load(f)
        if not rows:
            err(f"empty csv {os.path.relpath(f, OUT)}")


def check_insurance() -> None:
    m = os.path.join(OUT, "insurance_broking", "meridian_risk_partners")
    fk(f"{m}/02_policies_exposures/policies.csv", "client_id", f"{m}/01_clients_crm/clients.csv", "client_id")
    fk(f"{m}/02_policies_exposures/coverages.csv", "policy_id", f"{m}/02_policies_exposures/policies.csv", "policy_id")
    fk(f"{m}/03_marketing_submissions/submissions.csv", "client_id", f"{m}/01_clients_crm/clients.csv", "client_id")
    fk(f"{m}/03_marketing_submissions/quotes.csv", "submission_id", f"{m}/03_marketing_submissions/submissions.csv", "submission_id")
    fk(f"{m}/04_renewals/renewals.csv", "policy_id", f"{m}/02_policies_exposures/policies.csv", "policy_id")
    fk(f"{m}/05_certificates/certificate_requests.csv", "client_id", f"{m}/01_clients_crm/clients.csv", "client_id")
    fk(f"{m}/07_claims/claims.csv", "policy_id", f"{m}/02_policies_exposures/policies.csv", "policy_id")
    fk(f"{m}/08_billing_ar/invoices.csv", "policy_id", f"{m}/02_policies_exposures/policies.csv", "policy_id")
    fk(f"{m}/08_billing_ar/payments.csv", "invoice_id", f"{m}/08_billing_ar/invoices.csv", "invoice_id")
    fk(f"{m}/09_carrier_payables_commissions/producer_commissions.csv", "producer_id", f"{m}/12_hr_payroll/employees.csv", "employee_id")
    fk(f"{m}/11_finance_gl/general_ledger.csv", "account_number", f"{m}/11_finance_gl/chart_of_accounts.csv", "account_number")
    # planted: Meridian umbrella cert request exceeds limit
    cr = load(f"{m}/05_certificates/certificate_requests.csv")
    if not any(r["required_umbrella_limit"] == "5000000" for r in cr):
        err("Meridian: $5M umbrella cert request not found")
    # Castlebrook: workbook + dropped datasets
    c = os.path.join(OUT, "insurance_broking", "castlebrook_agency")
    if not os.path.exists(f"{c}/00_legacy_exports/MASTER_WORKBOOK_v7_FINAL (2).xlsx"):
        err("Castlebrook legacy workbook missing")
    if os.path.exists(f"{c}/14_workflow_events"):
        err("Castlebrook should have no workflow events")
    if os.path.exists(f"{c}/01_clients_crm/clients.csv"):
        err("Castlebrook clients.csv should be merged into workbook")


def check_industrial() -> None:
    n = os.path.join(OUT, "industrial_goods", "northfield_industrial_components")
    fk(f"{n}/02_sales_quote_to_order/sales_orders.csv", "customer_id", f"{n}/01_master_data/customers.csv", "customer_id")
    fk(f"{n}/02_sales_quote_to_order/sales_order_lines.csv", "so_number", f"{n}/02_sales_quote_to_order/sales_orders.csv", "so_number")
    fk(f"{n}/02_sales_quote_to_order/sales_order_lines.csv", "item_id", f"{n}/01_master_data/items.csv", "item_id")
    fk(f"{n}/03_procurement/purchase_orders.csv", "supplier_id", f"{n}/03_procurement/suppliers.csv", "supplier_id")
    fk(f"{n}/03_procurement/purchase_order_lines.csv", "po_number", f"{n}/03_procurement/purchase_orders.csv", "po_number")
    fk(f"{n}/04_receiving_ap/receipts.csv", "po_number", f"{n}/03_procurement/purchase_orders.csv", "po_number")
    fk(f"{n}/04_receiving_ap/supplier_invoices.csv", "po_number", f"{n}/03_procurement/purchase_orders.csv", "po_number")
    fk(f"{n}/06_warehouse_fulfillment/shipments.csv", "so_number", f"{n}/02_sales_quote_to_order/sales_orders.csv", "so_number")
    fk(f"{n}/11_billing_ar/customer_invoices.csv", "so_number", f"{n}/02_sales_quote_to_order/sales_orders.csv", "so_number")
    fk(f"{n}/11_billing_ar/customer_payments.csv", "applied_to_invoice", f"{n}/11_billing_ar/customer_invoices.csv", "invoice_number")
    fk(f"{n}/08_manufacturing/work_orders.csv", "item_id", f"{n}/01_master_data/items.csv", "item_id")
    fk(f"{n}/01_master_data/bills_of_material.csv", "component_item_id", f"{n}/01_master_data/items.csv", "item_id")
    fk(f"{n}/14_finance_gl/general_ledger.csv", "account_number", f"{n}/14_finance_gl/chart_of_accounts.csv", "account_number")
    fk(f"{n}/05_inventory/inventory_balances.csv", "item_id", f"{n}/01_master_data/items.csv", "item_id")
    # planted: three-way match price variance at Northfield
    twm = load(f"{n}/04_receiving_ap/three_way_match.csv")
    if not any(r["result"] == "Price variance" for r in twm):
        err("Northfield: no Price Variance in three_way_match")
    # planted: released ECO not in BOM
    ecos = load(f"{n}/13_engineering_compliance/engineering_change_orders.csv")
    if not any(r["status"] == "Released" and r["bom_updated"] == "N" for r in ecos):
        err("Northfield: released ECO with bom_updated=N not found")
    k = os.path.join(OUT, "industrial_goods", "keystone_bearing_and_drive")
    ws = openpyxl.load_workbook(f"{k}/05_inventory/inventory_balances.xlsx").active
    hdr = [c.value for c in ws[1]]
    if not any((row[hdr.index("On Hand Qty")] or 0) < 0 for row in ws.iter_rows(min_row=2, values_only=True)):
        err("Keystone: no negative inventory row")
    sups = load(f"{k}/03_procurement/suppliers.csv")
    if sum("grainger" in r["Supplier Name"].lower() for r in sups) < 2:
        err("Keystone: duplicate Grainger supplier not present")
    r = os.path.join(OUT, "industrial_goods", "ridgeway_fasteners_and_supply")
    if not os.path.exists(f"{r}/00_legacy_exports/RIDGEWAY MASTER (Fishbowl sync) v7 FINAL (2).xlsx"):
        err("Ridgeway legacy workbook missing")
    for should_not in ("17_workflow_events", "02_sales_quote_to_order", "05_inventory"):
        if os.path.exists(f"{r}/{should_not}"):
            err(f"Ridgeway should not have {should_not}")


def check_answer_key() -> None:
    items = json.load(open(os.path.join(OUT, "answer_key.json")))
    slug = {"Meridian": "insurance_broking/meridian_risk_partners", "Harborline": "insurance_broking/harborline_insurance_brokers",
            "Castlebrook": "insurance_broking/castlebrook_agency", "Northfield": "industrial_goods/northfield_industrial_components",
            "Keystone": "industrial_goods/keystone_bearing_and_drive", "Ridgeway": "industrial_goods/ridgeway_fasteners_and_supply"}
    ids = [i["id"] for i in items]
    if len(ids) != len(set(ids)):
        err(f"duplicate answer-key ids: {[i for i in ids if ids.count(i) > 1]}")
    for it in items:
        for ev in it["evidence"]:
            base = ev.split("#")[0]
            if base.endswith(")") and not base.endswith(".xlsx"):
                base = base.rsplit(" (", 1)[0]
            if "NOT AVAILABLE" in ev:
                continue
            if base.startswith("*/"):
                pats = [os.path.join(OUT, "*", "*", base[2:])]
            elif "/" in base and base.split("/")[0] in ("insurance_broking", "industrial_goods"):
                pats = [os.path.join(OUT, base)]
            elif base.split("/")[0] in {s.split("/")[1] for s in slug.values()}:
                pats = [os.path.join(OUT, "*", base)]
            else:
                pats = [os.path.join(OUT, slug[c], base) for c in it["companies"]]
            hits = [p for pat in pats for p in glob.glob(pat)] + [p for pat in pats for p in glob.glob(pat + ".*")]
            if not hits:
                err(f"{it['id']}: evidence path not found: {ev}")


def main() -> int:
    check_layout()
    check_insurance()
    check_industrial()
    check_answer_key()
    for e in errors:
        print("ERROR", e)
    print(f"{len(errors)} problems")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
